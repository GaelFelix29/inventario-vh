import io
import unittest

from openpyxl import Workbook

from services.importador_estados_cuenta import leer_estado_cuenta


def archivo_prueba(saldo_final=110):
    libro = Workbook()
    hoja = libro.active
    hoja["B2"] = "CUENTA: 1234"
    hoja["B3"] = "CLABE: 123456789012345678"
    hoja["B4"] = "EMPRESA DE PRUEBA"
    hoja["B5"] = "RFC: XAXX010101000"
    hoja["B9"] = "Fecha inicio: 01/08/2026"
    hoja["C9"] = "Fecha fin: 31/08/2026"
    for columna, valor in enumerate(["Fecha", "Descripción", "Referencia", "Cargo", "Abonos", "Saldo", "Clasificación"], 1):
        hoja.cell(10, columna, valor)
    for columna, valor in enumerate(["01/08/2026", "Depósito", "ABC", "", 10, saldo_final, "SPEI"], 1):
        hoja.cell(11, columna, valor)
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


class ImportadorEstadoCuentaTest(unittest.TestCase):
    def test_importa_y_concilia_movimientos(self):
        estado = leer_estado_cuenta(archivo_prueba(), "agosto.xlsx")
        self.assertEqual(estado["saldo_inicial"], 100)
        self.assertEqual(estado["saldo_final"], 110)
        self.assertEqual(len(estado["movimientos"]), 1)

    def test_rechaza_movimiento_sin_cargo_ni_abono(self):
        archivo = archivo_prueba()
        libro = __import__("openpyxl").load_workbook(archivo)
        hoja = libro.active
        hoja["E11"] = ""
        salida = io.BytesIO()
        libro.save(salida)
        salida.seek(0)
        with self.assertRaisesRegex(ValueError, "no contiene cargo ni abono"):
            leer_estado_cuenta(salida)


if __name__ == "__main__":
    unittest.main()
