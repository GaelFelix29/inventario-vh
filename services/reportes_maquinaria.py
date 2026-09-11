from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


VERDE, VERDE_OSCURO, VERDE_CLARO = "0B8F55", "07653E", "EAF7F0"


def _texto(valor, defecto="-"):
    return defecto if valor is None or str(valor).strip() in {"", "nan", "None"} else str(valor).strip()


def _texto_excel(valor, defecto="-"):
    texto = _texto(valor, defecto)
    return "'" + texto if texto.startswith(("=", "+", "-", "@")) else texto


def _tipo(maquina):
    valor = maquina.get("tipo") or {}
    return _texto(valor.get("nombre") if isinstance(valor, dict) else valor)


def _filtros_legibles(filtros):
    nombres = {"q": "Búsqueda", "tipo": "Tipo", "estado": "Estado", "ubicacion": "Ubicación"}
    activos = [f"{nombres[k]}: {v}" for k, v in filtros.items() if v]
    return activos or ["Sin filtros (inventario completo)"]


def filtrar_maquinarias(maquinas, filtros):
    terminos = (filtros.get("q") or "").casefold().split()
    tipo, estado, ubicacion = [(filtros.get(k) or "").casefold() for k in ("tipo", "estado", "ubicacion")]
    resultado = []
    for m in maquinas:
        tipo_m = _tipo(m)
        campos = [m.get(k) for k in ("id_activo", "descripcion", "categoria", "marca", "modelo", "numero_serie", "serie_interna", "ubicacion", "estado")] + [tipo_m]
        texto_busqueda = " ".join(_texto(v, "") for v in campos).casefold()
        if terminos and not all(termino in texto_busqueda for termino in terminos):
            continue
        if tipo and tipo != tipo_m.casefold(): continue
        if estado and estado != _texto(m.get("estado"), "").casefold(): continue
        if ubicacion and ubicacion != _texto(m.get("ubicacion"), "").casefold(): continue
        resultado.append(m)
    return resultado


def crear_excel_maquinaria(maquinas, filtros, generado_por):
    wb = Workbook(); ws = wb.active; ws.title = "Inventario"; ws.sheet_view.showGridLines = False; ws.freeze_panes = "A9"
    ws.merge_cells("A1:O1"); ws["A1"] = "VITAL HEALTH"; ws["A1"].font = Font(size=20, bold=True, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor=VERDE_OSCURO); ws["A1"].alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:O2"); ws["A2"] = "Reporte de Inventario de Maquinaria"; ws["A2"].font = Font(size=15, bold=True, color=VERDE_OSCURO)
    ws["A4"], ws["B4"] = "Generado:", datetime.now().strftime("%d/%m/%Y %H:%M"); ws["D4"], ws["E4"] = "Responsable:", generado_por; ws["H4"], ws["I4"] = "Resultados:", len(maquinas)
    for c in ("A4", "D4", "H4"): ws[c].font = Font(bold=True, color="66756D")
    ws.merge_cells("A6:O6"); ws["A6"] = "Filtros aplicados: " + " | ".join(_filtros_legibles(filtros)); ws["A6"].fill = PatternFill("solid", fgColor=VERDE_CLARO)
    ws["A6"].font = Font(bold=True, color=VERDE_OSCURO); ws["A6"].alignment = Alignment(vertical="center")
    headers = ["ID activo", "Descripción", "Categoría", "Marca", "Modelo", "Núm. de serie", "Serie interna", "Ubicación", "Tipo", "Estado", "Accesorios", "Precio unitario USD", "Total USD", "Valor MXN", "Fecha de alta"]
    for col, title in enumerate(headers, 1):
        c = ws.cell(8, col, title); c.fill = PatternFill("solid", fgColor=VERDE); c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(horizontal="center")
    borde = Border(bottom=Side(style="thin", color="DDE7E1"))
    for row, m in enumerate(maquinas, 9):
        values = [_texto_excel(m.get(k)) for k in ("id_activo", "descripcion", "categoria", "marca", "modelo", "numero_serie", "serie_interna", "ubicacion")] + [_texto_excel(_tipo(m)), _texto_excel(m.get("estado")), int(m.get("cantidad_accesorios") or 0), float(m.get("precio_unitario_us") or 0), float(m.get("total_us") or 0), float(m.get("valor_mx") or 0), _texto_excel(m.get("fecha_alta"))]
        for col, value in enumerate(values, 1):
            c = ws.cell(row, col, value); c.border = borde; c.alignment = Alignment(vertical="top", wrap_text=col in (2, 3)); c.fill = PatternFill("solid", fgColor="F5FAF7" if row % 2 == 0 else "FFFFFF")
        for col in (12, 13, 14): ws.cell(row, col).number_format = '$#,##0.00'
        ws.row_dimensions[row].height = 31
    last = 8 + max(len(maquinas), 1); ws.auto_filter.ref = f"A8:O{last}"
    for i, width in enumerate([15, 42, 27, 24, 19, 21, 18, 20, 15, 13, 13, 19, 17, 18, 16], 1): ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.print_title_rows = "1:8"
    output = BytesIO(); wb.save(output); output.seek(0); return output


def crear_pdf_maquinaria(maquinas, filtros, generado_por, logo_path=None):
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=12*mm, bottomMargin=14*mm, title="Reporte de Inventario de Maquinaria", author="Vital Health")
    styles = getSampleStyleSheet(); title = ParagraphStyle("vh-title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=20, textColor=colors.HexColor("#07653E")); cell = ParagraphStyle("vh-cell", parent=styles["BodyText"], fontSize=6.2, leading=7.4, textColor=colors.HexColor("#26382F")); small = ParagraphStyle("vh-small", parent=cell, fontSize=7, leading=9)
    logo = Image(str(Path(logo_path)), width=31*mm, height=15*mm) if logo_path and Path(logo_path).exists() else Paragraph("<b>VITAL HEALTH</b>", title)
    head = Table([[logo, Paragraph("<b>Reporte de Inventario de Maquinaria</b><br/><font size='8'>Control y seguimiento de activos</font>", title), Paragraph(f"<b>{len(maquinas)} activos</b><br/>{datetime.now().strftime('%d/%m/%Y %H:%M')}<br/>{_texto(generado_por)}", small)]], colWidths=[38*mm, 165*mm, 65*mm]); head.setStyle(TableStyle([("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("ALIGN", (-1,0), (-1,0), "RIGHT")]))
    filter_box = Table([[Paragraph("<b>Filtros aplicados:</b> " + " &nbsp; | &nbsp; ".join(_filtros_legibles(filtros)), small)]], colWidths=[268*mm]); filter_box.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#EAF7F0")), ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#C9E7D7")), ("PADDING", (0,0), (-1,-1), 6)]))
    headers = ["ID", "Descripción", "Categoría", "Marca / modelo", "Serie", "Ubicación", "Tipo", "Estado", "Acc.", "Valor MXN"]
    data = [[Paragraph(f"<b>{h}</b>", cell) for h in headers]]
    for m in maquinas:
        vals = [escape(_texto(m.get("id_activo"))), escape(_texto(m.get("descripcion"))), escape(_texto(m.get("categoria"))), f"{escape(_texto(m.get('marca')))}<br/>{escape(_texto(m.get('modelo')))}", escape(_texto(m.get("numero_serie"))), escape(_texto(m.get("ubicacion"))), escape(_tipo(m)), escape(_texto(m.get("estado"))), str(int(m.get("cantidad_accesorios") or 0)), f"${float(m.get('valor_mx') or 0):,.2f}"]
        data.append([Paragraph(v, cell) for v in vals])
    if not maquinas: data.append([Paragraph("Sin resultados", cell)] + [""]*9)
    table = Table(data, repeatRows=1, colWidths=[18*mm, 50*mm, 35*mm, 35*mm, 28*mm, 25*mm, 18*mm, 17*mm, 10*mm, 30*mm]); table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0B8F55")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("VALIGN", (0,0), (-1,-1), "TOP"), ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#DDE7E1")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F5FAF7")]), ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4), ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)]))
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#607168")); canvas.drawString(10*mm, 6*mm, "Vital Health - Sistema de Inventario de Activos"); canvas.drawRightString(287*mm, 6*mm, f"Página {document.page}"); canvas.restoreState()
    doc.build([head, Spacer(1, 4*mm), filter_box, Spacer(1, 5*mm), table], onFirstPage=footer, onLaterPages=footer); output.seek(0); return output
