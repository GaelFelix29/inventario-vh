from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

VERDE = "078C56"
VERDE_OSCURO = "075C3B"


def _texto_filtros(filtros):
    partes = []
    if filtros.get("buscar"):
        partes.append(f"Búsqueda: {filtros['buscar']}")
    if filtros.get("tipo") != "todos":
        partes.append(f"Tipo: {filtros['tipo'].title()}")
    concepto = filtros.get("concepto", "TODOS")
    if concepto != "TODOS":
        etiqueta = "PAGOSA (Asociado)" if concepto == "PAGOSA" else "PAGOS (Comisionista PF)"
        partes.append(f"Concepto: {etiqueta}")
    if filtros.get("fecha_desde"):
        partes.append(f"Desde: {filtros['fecha_desde']:%d/%m/%Y}")
    if filtros.get("fecha_hasta"):
        partes.append(f"Hasta: {filtros['fecha_hasta']:%d/%m/%Y}")
    return partes or ["Sin filtros adicionales"]


def crear_excel_movimientos(estado, movimientos, filtros, generado_por):
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A9"
    ws.merge_cells("A1:F1")
    ws["A1"] = "VITAL HEALTH · MOVIMIENTOS BANCARIOS"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=VERDE_OSCURO)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:F2")
    ws["A2"] = f"{estado.banco} · •••• {estado.numero_cuenta[-4:]} · {estado.periodo_inicio:%d/%m/%Y} al {estado.periodo_fin:%d/%m/%Y}"
    ws["A2"].font = Font(bold=True, color="355A48")
    ws.merge_cells("A3:F3")
    ws["A3"] = " | ".join(_texto_filtros(filtros))
    ws["A3"].font = Font(italic=True, color="687A70")
    ws.merge_cells("A4:F4")
    ws["A4"] = f"Generado por {generado_por} · {datetime.now():%d/%m/%Y %H:%M} · {len(movimientos):,} registros"
    ws["A4"].font = Font(size=9, color="7D8B84")
    resumen = [("Saldo inicial", estado.saldo_inicial), ("Cargos", estado.total_cargos),
               ("Abonos", estado.total_abonos), ("Saldo final", estado.saldo_final)]
    for columna, (titulo, valor) in enumerate(resumen, 1):
        celda = ws.cell(6, columna, titulo)
        celda.font = Font(size=9, color="6C7D73")
        monto = ws.cell(7, columna, float(valor or 0))
        monto.number_format = '$#,##0.00;[Red]-$#,##0.00'
        monto.font = Font(size=12, bold=True, color=VERDE if titulo == "Abonos" else "C63F4D" if titulo == "Cargos" else "17251E")
    encabezados = ["Fecha", "Descripción", "Referencia", "Cargo", "Abono", "Saldo"]
    for columna, titulo in enumerate(encabezados, 1):
        celda = ws.cell(9, columna, titulo)
        celda.fill = PatternFill("solid", fgColor=VERDE)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center")
    borde = Border(bottom=Side(style="thin", color="DCE8E1"))
    for indice, movimiento in enumerate(movimientos, 10):
        valores = [movimiento.fecha, movimiento.descripcion, movimiento.referencia,
                   movimiento.cargo, movimiento.abono, movimiento.saldo]
        for columna, valor in enumerate(valores, 1):
            celda = ws.cell(indice, columna, valor)
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=columna in (2, 3))
        ws.cell(indice, 1).number_format = "dd/mm/yyyy"
        for columna in (4, 5, 6):
            ws.cell(indice, columna).number_format = '$#,##0.00;[Red]-$#,##0.00'
    anchos = [14, 66, 34, 18, 18, 18]
    for columna, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(columna)].width = ancho
    ws.auto_filter.ref = f"A9:F{max(9, 9 + len(movimientos))}"
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida


def crear_pdf_movimientos(estado, movimientos, filtros, generado_por, logo_path):
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), rightMargin=10*mm,
                            leftMargin=10*mm, topMargin=11*mm, bottomMargin=12*mm,
                            title="Movimientos bancarios", author="Vital Health")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("titulo", parent=estilos["Heading1"], fontSize=16,
                            textColor=colors.HexColor("#075C3B"), leading=19)
    texto = ParagraphStyle("texto", parent=estilos["BodyText"], fontSize=7.5,
                           leading=10, textColor=colors.HexColor("#34483E"))
    derecha = ParagraphStyle("derecha", parent=texto, alignment=TA_RIGHT)
    logo = Image(logo_path, 28*mm, 18*mm) if Path(logo_path).exists() else ""
    cabecera = Table([[logo, Paragraph(
        f"<b>Movimientos bancarios</b><br/><font size='9'>{estado.banco} · •••• {estado.numero_cuenta[-4:]} · {estado.periodo_inicio:%d/%m/%Y} al {estado.periodo_fin:%d/%m/%Y}</font>", titulo),
        Paragraph(f"<b>{len(movimientos):,} registros</b><br/>{datetime.now():%d/%m/%Y %H:%M}<br/>{generado_por}", derecha)]],
        colWidths=[34*mm, 175*mm, 65*mm])
    cabecera.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    filtros_tabla = Table([[Paragraph("<b>Filtros:</b> " + " | ".join(_texto_filtros(filtros)), texto)]], colWidths=[274*mm])
    filtros_tabla.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#EAF7F0")), ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#C9E7D7")), ("PADDING", (0,0), (-1,-1), 6)]))
    resumen_estilo = ParagraphStyle(
        "resumen", parent=texto, fontSize=8, leading=11,
        textColor=colors.HexColor("#687A70")
    )
    resumen = Table([[
        Paragraph(
            "Saldo inicial<br/><font size='13' color='#17251E'><b>"
            f"${float(estado.saldo_inicial or 0):,.2f}</b></font>",
            resumen_estilo,
        ),
        Paragraph(
            "Cargos<br/><font size='13' color='#C63F4D'><b>"
            f"${float(estado.total_cargos or 0):,.2f}</b></font>",
            resumen_estilo,
        ),
        Paragraph(
            "Abonos<br/><font size='13' color='#078C56'><b>"
            f"${float(estado.total_abonos or 0):,.2f}</b></font>",
            resumen_estilo,
        ),
    ]], colWidths=[89*mm, 89*mm, 89*mm], hAlign="LEFT")
    resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#DCE8E1")),
        ("INNERGRID", (0, 0), (-1, -1), .5, colors.HexColor("#DCE8E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    datos = [["Fecha", "Descripción", "Referencia", "Cargo", "Abono", "Saldo"]]
    for m in movimientos:
        datos.append([m.fecha.strftime("%d/%m/%Y"), Paragraph(str(m.descripcion or ""), texto),
                      Paragraph(str(m.referencia or ""), texto),
                      f"${float(m.cargo or 0):,.2f}" if m.cargo else "—",
                      f"${float(m.abono or 0):,.2f}" if m.abono else "—",
                      f"${float(m.saldo or 0):,.2f}"])
    tabla = Table(datos, repeatRows=1, colWidths=[22*mm, 104*mm, 57*mm, 29*mm, 29*mm, 33*mm])
    tabla.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#078C56")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 7), ("ALIGN", (3,1), (-1,-1), "RIGHT"), ("VALIGN", (0,0), (-1,-1), "TOP"), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F4F9F6")]), ("GRID", (0,0), (-1,-1), .3, colors.HexColor("#DCE8E1")), ("PADDING", (0,0), (-1,-1), 4)]))
    doc.build([
        cabecera, Spacer(1, 5*mm), filtros_tabla, Spacer(1, 3*mm),
        resumen, Spacer(1, 4*mm), tabla
    ])
    salida.seek(0)
    return salida
