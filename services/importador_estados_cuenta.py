import hashlib
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


CENTAVO = Decimal("0.01")
ENCABEZADOS = ["fecha", "descripcion", "referencia", "cargo", "abonos", "saldo", "clasificacion"]


def _texto(valor):
    return str(valor or "").strip()


def _decimal(valor):
    if valor in (None, ""):
        return Decimal("0.00")
    try:
        return Decimal(str(valor).replace(",", "")).quantize(CENTAVO)
    except InvalidOperation as error:
        raise ValueError(f"Importe inválido: {valor}") from error


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(_texto(valor), formato).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha inválida: {valor}")


def _normalizar_encabezado(valor):
    tabla = str.maketrans("áéíóúñ", "aeioun")
    return _texto(valor).lower().translate(tabla).replace(" ", "")


def leer_estado_cuenta(contenido, nombre_archivo="estado_cuenta.xlsx"):
    """Lee y concilia el formato de estado de cuenta Banregio."""
    datos = contenido if hasattr(contenido, "read") else io.BytesIO(contenido)
    if hasattr(datos, "seek"):
        datos.seek(0)
    bytes_archivo = datos.read()
    libro = load_workbook(io.BytesIO(bytes_archivo), read_only=True, data_only=True)
    hoja = libro.active

    encabezados = [_normalizar_encabezado(hoja.cell(10, col).value) for col in range(1, 8)]
    encabezados[1] = "descripcion" if encabezados[1].startswith("descripci") else encabezados[1]
    encabezados[6] = "clasificacion" if encabezados[6].startswith("clasificaci") else encabezados[6]
    if encabezados != ENCABEZADOS:
        raise ValueError("El archivo no tiene el formato esperado de Banregio.")

    cuenta = _texto(hoja["B2"].value).removeprefix("CUENTA:").strip()
    clabe = _texto(hoja["B3"].value).removeprefix("CLABE:").strip()
    titular = _texto(hoja["B4"].value)
    rfc = _texto(hoja["B5"].value).removeprefix("RFC:").strip()
    periodo_inicio = _fecha(_texto(hoja["B9"].value).removeprefix("Fecha inicio:").strip())
    periodo_fin = _fecha(_texto(hoja["C9"].value).removeprefix("Fecha fin:").strip())
    if not cuenta or not clabe:
        raise ValueError("No se encontró la cuenta o la CLABE del estado de cuenta.")

    movimientos = []
    saldo_anterior = None
    total_cargos = Decimal("0.00")
    total_abonos = Decimal("0.00")
    for renglon, valores in enumerate(hoja.iter_rows(min_row=11, max_col=7, values_only=True), start=11):
        if not any(valor not in (None, "") for valor in valores):
            continue
        fecha, descripcion, referencia, cargo, abono, saldo, clasificacion = valores
        cargo, abono, saldo = _decimal(cargo), _decimal(abono), _decimal(saldo)
        if cargo > 0 and abono > 0:
            raise ValueError(f"El renglón {renglon} contiene cargo y abono simultáneamente.")
        if cargo == 0 and abono == 0:
            raise ValueError(f"El renglón {renglon} no contiene cargo ni abono.")
        fecha = _fecha(fecha)
        if fecha < periodo_inicio or fecha > periodo_fin:
            raise ValueError(f"El renglón {renglon} está fuera del periodo declarado.")
        saldo_inicial_fila = saldo + cargo - abono
        if saldo_anterior is not None and saldo_inicial_fila != saldo_anterior:
            raise ValueError(f"El saldo no concilia en el renglón {renglon}.")
        huella_base = "|".join(map(str, (renglon, fecha, _texto(descripcion), _texto(referencia), cargo, abono, saldo)))
        movimientos.append({
            "renglon_origen": renglon, "fecha": fecha, "descripcion": _texto(descripcion),
            "referencia": _texto(referencia), "cargo": cargo, "abono": abono,
            "saldo": saldo, "clasificacion_banco": _texto(clasificacion) or None,
            "huella": hashlib.sha256(huella_base.encode("utf-8")).hexdigest(),
        })
        total_cargos += cargo
        total_abonos += abono
        saldo_anterior = saldo

    if not movimientos:
        raise ValueError("El estado de cuenta no contiene movimientos.")
    saldo_inicial = movimientos[0]["saldo"] + movimientos[0]["cargo"] - movimientos[0]["abono"]
    saldo_final = movimientos[-1]["saldo"]
    if saldo_inicial - total_cargos + total_abonos != saldo_final:
        raise ValueError("Los totales del archivo no concilian con el saldo final.")

    return {
        "banco": "BANREGIO", "numero_cuenta": cuenta, "clabe": clabe,
        "titular": titular, "rfc": rfc, "periodo_inicio": periodo_inicio,
        "periodo_fin": periodo_fin, "saldo_inicial": saldo_inicial,
        "saldo_final": saldo_final, "total_cargos": total_cargos,
        "total_abonos": total_abonos, "nombre_archivo": nombre_archivo,
        "archivo_hash": hashlib.sha256(bytes_archivo).hexdigest(),
        "movimientos": movimientos,
    }
